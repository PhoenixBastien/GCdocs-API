from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from config import gcdocs

# node_id = 42279688
node_id = 41737208
node = gcdocs.get_node(node_id=node_id)
filename = gcdocs.get_result_value(response=node, key="name")
extension = Path(filename).suffix
print(extension)

content = gcdocs.get_document_content(node_id=node_id)
stream = BytesIO(content)

with open(f"out/{filename} ({node_id}).txt", "w") as f:
    if extension.startswith(".doc"):
        doc = Document(stream)

        for para in doc.paragraphs:
            f.write(para.text + "\n")
    elif extension.startswith(".xls"):
        wb = load_workbook(stream, data_only=True)

        for cell in wb.worksheets[0]:
            print(cell)

stream.close()
